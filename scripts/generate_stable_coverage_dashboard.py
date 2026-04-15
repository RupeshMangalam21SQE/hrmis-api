import json
import re
import csv
import ast
import argparse
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "tests" / "test_data" / "HRMIS_API-Test cases .xlsx"
CONFIG_PATH = ROOT / "docs" / "coverage" / "stable_modules.json"
VERIFIED_MAP_PATH = ROOT / "docs" / "coverage" / "verified_case_mapping.json"
OUT_DIR = ROOT / "docs" / "coverage"
ENDPOINTS_DIR = ROOT / "src" / "endpoints"
CLIENTS_DIR = ROOT / "src" / "clients"
CONTRACTS_DIR = ROOT / "tests" / "contracts"
REPORTS_DIR = ROOT / "reports"
CASE_ID_PATTERNS = (
    re.compile(r"^[A-Z]{2,5}-[A-Z]{1,3}-\d+$"),  # AUTH-PF-001, SAL-UP-001
    re.compile(r"^[A-Z]{2,5}-\d+$"),  # ANN-001
    re.compile(r"^TC_[A-Z0-9]+_\d+$"),  # TC_GET_01, TC_OB_01
)


def is_case_id(value: str) -> bool:
    return any(p.match(value) for p in CASE_ID_PATTERNS)


def _cell(v):
    return "" if v is None else str(v).strip()


def normalize_endpoint(path: str) -> str:
    if not path:
        return ""
    p = path.strip()
    if p.startswith("http://") or p.startswith("https://"):
        parts = p.split("/", 3)
        p = "/" + parts[3] if len(parts) > 3 else "/"
    p = p.split("?", 1)[0].strip()
    if not p:
        return ""
    if not p.startswith("/"):
        p = "/" + p
    p = re.sub(r"/+", "/", p)
    p = p.rstrip("/") if p != "/" else p
    api_prefix = "/HRMBackendTest"
    if p.startswith(api_prefix + "/"):
        p = p[len(api_prefix) :]
    return p


def split_endpoints(raw: str) -> list[str]:
    value = _cell(raw).replace("\r", "\n")
    out = []
    for part in value.split("\n"):
        s = part.strip()
        if not s:
            continue
        if " " in s and s.startswith("/"):
            s = s.split(" ", 1)[0]
        out.append(s)
    return out


def load_config():
    return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))


def load_verified_mapping():
    if not VERIFIED_MAP_PATH.exists():
        return {"module_case_mapping": {}}
    return json.loads(VERIFIED_MAP_PATH.read_text(encoding="utf-8"))


def _parse_verified_id_set(module_entry: dict, key: str) -> set[str]:
    raw = (module_entry or {}).get(key, [])
    if isinstance(raw, list):
        return {str(v).strip() for v in raw if str(v).strip()}
    if isinstance(raw, dict):
        return {str(k).strip() for k in raw.keys() if str(k).strip()}
    return set()


def _get_case_evidence(module_entry: dict, bucket: str, case_id: str) -> dict:
    evidence = ((module_entry or {}).get("evidence", {}) or {}).get(bucket, {})
    if not isinstance(evidence, dict):
        return {"files": [], "note": ""}
    row = evidence.get(case_id, {})
    if not isinstance(row, dict):
        return {"files": [], "note": ""}
    files = row.get("files", [])
    if not isinstance(files, list):
        files = []
    note = str(row.get("note", "")).strip()
    return {"files": [str(x).strip() for x in files if str(x).strip()], "note": note}


def _module_hints(module_key: str) -> list[str]:
    hints = {
        "login": ["login", "auth"],
        "salary_management": ["salary_management", "salary"],
        "announcements": ["announcements", "announcement"],
        "dashboard": ["dashboard"],
        "onboarding": ["onboarding"],
    }
    return hints.get(module_key, [module_key])


def find_latest_junit_xml() -> Path | None:
    if not REPORTS_DIR.exists():
        return None
    candidates = [p for p in REPORTS_DIR.rglob("junit.xml") if p.is_file()]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def load_junit_module_counts(config: dict) -> tuple[dict[str, dict[str, int]], str | None]:
    junit_path = find_latest_junit_xml()
    if not junit_path:
        return {}, None

    counts = {
        m["module_key"]: {"passed": 0, "failed": 0, "skipped": 0, "total": 0}
        for m in config["modules"]
    }

    try:
        root = ET.parse(junit_path).getroot()
    except ET.ParseError:
        return counts, str(junit_path.relative_to(ROOT)).replace("\\", "/")

    for tc in root.findall(".//testcase"):
        classname = (tc.attrib.get("classname") or "").lower()
        matched_module = None
        for m in config["modules"]:
            mk = m["module_key"]
            if any(h in classname for h in _module_hints(mk)):
                matched_module = mk
                break

        if not matched_module:
            continue

        c = counts[matched_module]
        c["total"] += 1
        has_failure = tc.find("failure") is not None or tc.find("error") is not None
        skipped_node = tc.find("skipped")

        if has_failure:
            c["failed"] += 1
            continue

        if skipped_node is not None:
            s_type = (skipped_node.attrib.get("type") or "").lower()
            s_msg = (skipped_node.attrib.get("message") or "").lower()
            # Treat xfail as failed for pass-vs-fail visibility.
            if "xfail" in s_type or "xfail" in s_msg:
                c["failed"] += 1
            else:
                c["skipped"] += 1
            continue

        c["passed"] += 1

    return counts, str(junit_path.relative_to(ROOT)).replace("\\", "/")


def load_manual_cases_by_sheet(workbook_path):
    wb = load_workbook(workbook_path, data_only=True)
    result = {}
    for ws in wb.worksheets:
        rows = []
        header = [_cell(v).lower() for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        idx_method = header.index("http method") if "http method" in header else None
        idx_endpoint = header.index("api endpoint") if "api endpoint" in header else None

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            found = None
            for col_val in row:
                value = _cell(col_val)
                if is_case_id(value):
                    found = value
                    break
            if found:
                method = _cell(row[idx_method]) if idx_method is not None and idx_method < len(row) else ""
                endpoint_raw = _cell(row[idx_endpoint]) if idx_endpoint is not None and idx_endpoint < len(row) else ""
                endpoint_variants = split_endpoints(endpoint_raw)
                endpoint_norm = normalize_endpoint(endpoint_variants[0]) if endpoint_variants else ""
                rows.append(
                    {
                        "case_id": found,
                        "http_method": method,
                        "api_endpoint": endpoint_raw,
                        "endpoint_norm": endpoint_norm,
                    }
                )

        dedup = {}
        for r in rows:
            dedup[r["case_id"]] = r
        result[ws.title] = sorted(dedup.values(), key=lambda x: x["case_id"])
    return result


def extract_case_ids_from_tests(files):
    ids = set()
    for f in files:
        text = f.read_text(encoding="utf-8")
        ids.update(re.findall(r"\b[A-Z]{2,5}-[A-Z]{1,3}-\d+\b", text))
        ids.update(re.findall(r"\b[A-Z]{2,5}-\d+\b", text))
        ids.update(re.findall(r"\bTC_[A-Z0-9]+_\d+\b", text))
    return ids


def load_endpoint_constants() -> dict[str, str]:
    constants: dict[str, str] = {}
    for f in ENDPOINTS_DIR.glob("*.py"):
        if f.name.startswith("__"):
            continue
        src = f.read_text(encoding="utf-8")
        try:
            tree = ast.parse(src)
        except SyntaxError:
            continue
        for node in tree.body:
            if isinstance(node, ast.Assign) and isinstance(node.value, ast.Constant) and isinstance(node.value.value, str):
                for t in node.targets:
                    if isinstance(t, ast.Name) and t.id.isupper():
                        constants[t.id] = normalize_endpoint(node.value.value)
    return constants


def client_method_endpoint_map(endpoint_constants: dict[str, str]) -> dict[tuple[str, str], set[str]]:
    method_map: dict[tuple[str, str], set[str]] = {}
    for f in CLIENTS_DIR.glob("*.py"):
        if f.name.startswith("__"):
            continue
        src = f.read_text(encoding="utf-8")
        try:
            tree = ast.parse(src)
        except SyntaxError:
            continue

        class_defs = [n for n in tree.body if isinstance(n, ast.ClassDef) and n.name.endswith("Client")]
        for c in class_defs:
            for fn in [n for n in c.body if isinstance(n, ast.FunctionDef)]:
                paths: set[str] = set()
                for n in ast.walk(fn):
                    if isinstance(n, ast.Name) and n.id in endpoint_constants:
                        paths.add(endpoint_constants[n.id])
                    if isinstance(n, ast.Constant) and isinstance(n.value, str):
                        s = n.value.strip()
                        if "/" in s and (
                            s.startswith("/")
                            or s.startswith("api/")
                            or s.startswith("menu/")
                            or s.startswith("salary/")
                            or s.startswith("assest/")
                            or s.startswith("onboarding/")
                            or s.startswith("document/")
                            or s.startswith("announcement")
                        ):
                            paths.add(normalize_endpoint(s))
                method_map[(c.name, fn.name)] = {p for p in paths if p}
    return method_map


def extract_contract_spec_paths(endpoint_constants: dict[str, str], client_method_map: dict[tuple[str, str], set[str]]) -> set[str]:
    paths: set[str] = set()
    for f in CONTRACTS_DIR.glob("test_*.py"):
        src = f.read_text(encoding="utf-8")
        try:
            tree = ast.parse(src)
        except SyntaxError:
            continue

        # direct constant names and literal endpoints
        for n in ast.walk(tree):
            if isinstance(n, ast.Name) and n.id in endpoint_constants:
                paths.add(endpoint_constants[n.id])
            if isinstance(n, ast.Constant) and isinstance(n.value, str):
                s = n.value.strip()
                if "/" in s and (
                    s.startswith("/")
                    or s.startswith("api/")
                    or s.startswith("menu/")
                    or s.startswith("salary/")
                    or s.startswith("assest/")
                    or s.startswith("onboarding/")
                    or s.startswith("document/")
                    or s.startswith("announcement")
                ):
                    paths.add(normalize_endpoint(s))

        # infer client calls in each function
        for fn in [n for n in ast.walk(tree) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))]:
            var_class: dict[str, str] = {}
            for n in ast.walk(fn):
                if isinstance(n, ast.Assign) and isinstance(n.value, ast.Call) and isinstance(n.value.func, ast.Name):
                    cls = n.value.func.id
                    if cls.endswith("Client"):
                        for t in n.targets:
                            if isinstance(t, ast.Name):
                                var_class[t.id] = cls
                if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and isinstance(n.func.value, ast.Name):
                    var = n.func.value.id
                    method = n.func.attr
                    cls = var_class.get(var)
                    if cls:
                        paths.update(client_method_map.get((cls, method), set()))
    return {p for p in paths if p}


def schema_file_count(schema_dir):
    p = ROOT / schema_dir
    if not p.exists() or not p.is_dir():
        return 0
    return len(list(p.glob("*.json")))


def pct(a, b):
    if b == 0:
        return 0.0
    return round((a / b) * 100.0, 2)


def main(strict_verified_spec: bool = False):
    config = load_config()
    verified = load_verified_mapping().get("module_case_mapping", {})
    manual = load_manual_cases_by_sheet(EXCEL_PATH)
    junit_counts, junit_source = load_junit_module_counts(config)
    endpoint_constants = load_endpoint_constants()
    client_method_map = client_method_endpoint_map(endpoint_constants)
    spec_paths = extract_contract_spec_paths(endpoint_constants, client_method_map)

    rows = []
    trace_rows = []
    strict_violations = []
    totals = {
        "manual_cases": 0,
        "automated_cases": 0,
        "spec_flow_cases": 0,
        "overlap_cases": 0,
        "net_covered_cases": 0,
        "pending_cases": 0,
        "go_signal_modules": 0,
        "module_count": 0,
        "schema_files": 0,
        "passed_cases": 0,
        "failed_cases": 0,
    }

    for m in config["modules"]:
        module_key = m["module_key"]
        sheet = m["sheet_name"]
        go_signal = bool(m.get("go_signal", False))

        manual_rows = manual.get(sheet, [])
        manual_ids = {r["case_id"] for r in manual_rows}

        test_files = sorted(ROOT.glob(m["test_glob"]))
        auto_ids_all = extract_case_ids_from_tests(test_files)
        module_verified = verified.get(module_key, {}) or {}
        verified_auto = _parse_verified_id_set(module_verified, "automated_case_ids")
        verified_spec = _parse_verified_id_set(module_verified, "spec_case_ids")

        auto_ids = set(manual_ids.intersection(auto_ids_all.union(verified_auto)))
        spec_ids = set(manual_ids.intersection(verified_spec))

        overlap = auto_ids.intersection(spec_ids)
        net_covered = auto_ids.union(spec_ids)
        pending = manual_ids - net_covered

        schema_count = schema_file_count(m["schema_dir"])
        run_counts = junit_counts.get(module_key, {})
        passed_cases = run_counts.get("passed", 0)
        failed_cases = run_counts.get("failed", 0)

        auto_evidence_present = sum(
            1 for cid in auto_ids if _get_case_evidence(module_verified, "automated", cid)["files"]
        )
        spec_evidence_present = sum(
            1 for cid in spec_ids if _get_case_evidence(module_verified, "spec", cid)["files"]
        )

        row = {
            "module_key": module_key,
            "sheet_name": sheet,
            "go_signal": go_signal,
            "manual_cases": len(manual_ids),
            "automated_cases": len(auto_ids),
            "spec_flow_cases": len(spec_ids),
            "overlap_cases": len(overlap),
            "net_covered_cases": len(net_covered),
            "pending_cases": len(pending),
            "coverage_percent": pct(len(net_covered), len(manual_ids)),
            "test_files": len(test_files),
            "schema_files": schema_count,
            "passed_cases": passed_cases,
            "failed_cases": failed_cases,
            "auto_evidence_present": auto_evidence_present,
            "auto_evidence_missing": max(0, len(auto_ids) - auto_evidence_present),
            "spec_evidence_present": spec_evidence_present,
            "spec_evidence_missing": max(0, len(spec_ids) - spec_evidence_present),
        }
        rows.append(row)

        for r in manual_rows:
            cid = r["case_id"]
            auto_flag = cid in auto_ids
            spec_flag = cid in spec_ids
            if strict_verified_spec and spec_flag and cid not in verified_spec:
                strict_violations.append(
                    f"{module_key}:{cid} is marked spec-covered but is missing from verified spec mapping"
                )

            auto_evidence = _get_case_evidence(module_verified, "automated", cid) if auto_flag else {"files": [], "note": ""}
            spec_evidence = _get_case_evidence(module_verified, "spec", cid) if spec_flag else {"files": [], "note": ""}
            if auto_flag and spec_flag:
                bucket = "automated+spec"
            elif auto_flag:
                bucket = "automated_only"
            elif spec_flag:
                bucket = "spec_only"
            else:
                bucket = "uncovered"
            trace_rows.append(
                {
                    "module_key": module_key,
                    "sheet_name": sheet,
                    "case_id": cid,
                    "http_method": r["http_method"],
                    "api_endpoint": r["api_endpoint"],
                    "normalized_endpoint": r["endpoint_norm"],
                    "automated_covered": str(auto_flag),
                    "spec_flow_covered": str(spec_flag),
                    "overlap": str(auto_flag and spec_flag),
                    "coverage_bucket": bucket,
                    "automated_evidence_files": "; ".join(auto_evidence["files"]),
                    "spec_evidence_files": "; ".join(spec_evidence["files"]),
                    "evidence_note": spec_evidence["note"] or auto_evidence["note"],
                }
            )

        totals["module_count"] += 1
        totals["schema_files"] += schema_count
        if go_signal:
            totals["go_signal_modules"] += 1
            totals["manual_cases"] += len(manual_ids)
            totals["automated_cases"] += len(auto_ids)
            totals["spec_flow_cases"] += len(spec_ids)
            totals["overlap_cases"] += len(overlap)
            totals["net_covered_cases"] += len(net_covered)
            totals["pending_cases"] += len(pending)
            totals["passed_cases"] += passed_cases
            totals["failed_cases"] += failed_cases

    totals["coverage_percent"] = pct(totals["net_covered_cases"], totals["manual_cases"])

    if strict_verified_spec and strict_violations:
        for v in strict_violations:
            print(f"STRICT-ERROR: {v}")
        raise SystemExit(
            f"Strict verified-spec mode failed with {len(strict_violations)} violation(s)."
        )

    payload = {
        "generated_at_utc": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
        "excel_path": str(EXCEL_PATH.relative_to(ROOT)).replace("\\", "/"),
        "config_path": str(CONFIG_PATH.relative_to(ROOT)).replace("\\", "/"),
        "verified_mapping_path": str(VERIFIED_MAP_PATH.relative_to(ROOT)).replace("\\", "/"),
        "strict_verified_spec": strict_verified_spec,
        "junit_source": junit_source,
        "contract_spec_paths_total": len(spec_paths),
        "summary": totals,
        "modules": rows,
    }

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "stable_coverage_dashboard.json").write_text(
        json.dumps(payload, indent=2), encoding="utf-8"
    )

    with (OUT_DIR / "case_traceability.csv").open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "module_key",
                "sheet_name",
                "case_id",
                "http_method",
                "api_endpoint",
                "normalized_endpoint",
                "automated_covered",
                "spec_flow_covered",
                "overlap",
                "coverage_bucket",
                "automated_evidence_files",
                "spec_evidence_files",
                "evidence_note",
            ],
        )
        writer.writeheader()
        writer.writerows(trace_rows)

    (OUT_DIR / "contract_spec_paths.txt").write_text(
        "\n".join(sorted(spec_paths)) + ("\n" if spec_paths else ""), encoding="utf-8"
    )

    lines = []
    lines.append("# HRMS Stable API Coverage Dashboard")
    lines.append("")
    lines.append(f"Generated At (UTC): {payload['generated_at_utc']}")
    lines.append(f"Manual Test Source: {payload['excel_path']}")
    lines.append(f"Stability Config: {payload['config_path']}")
    lines.append(f"Verified Mapping: {payload['verified_mapping_path']}")
    lines.append(f"JUnit Source: {payload['junit_source'] or 'not found'}")
    lines.append("")
    lines.append("## Go-Signal Summary")
    lines.append(f"- Modules with go signal: {totals['go_signal_modules']}/{totals['module_count']}")
    lines.append(f"- Manual cases (go-signal modules): {totals['manual_cases']}")
    lines.append(f"- Automated cases mapped to manual IDs: {totals['automated_cases']}")
    lines.append(f"- Spec-flow mapped cases: {totals['spec_flow_cases']}")
    lines.append(f"- Overlap (automated + spec): {totals['overlap_cases']}")
    lines.append(f"- Net covered (de-duplicated): {totals['net_covered_cases']}")
    lines.append(f"- Pending (go-signal modules): {totals['pending_cases']}")
    lines.append(f"- Passed test cases (go-signal modules): {totals['passed_cases']}")
    lines.append(f"- Failed test cases (go-signal modules): {totals['failed_cases']}")
    lines.append(f"- Coverage (go-signal modules): {totals['coverage_percent']}%")
    lines.append(f"- Schema files across configured modules: {totals['schema_files']}")
    lines.append(f"- Contract/spec paths inferred: {len(spec_paths)}")
    lines.append("")
    lines.append("## Module Breakdown")
    lines.append("| Module | Go Signal | Manual | Automated | Spec Flow | Overlap | Net Covered | Pending | Passed | Failed | Coverage % | Test Files | Schema Files |")
    lines.append("|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|")
    for r in rows:
        lines.append(
            f"| {r['module_key']} | {'Yes' if r['go_signal'] else 'No'} | {r['manual_cases']} | {r['automated_cases']} | {r['spec_flow_cases']} | {r['overlap_cases']} | {r['net_covered_cases']} | {r['pending_cases']} | {r['passed_cases']} | {r['failed_cases']} | {r['coverage_percent']} | {r['test_files']} | {r['schema_files']} |"
        )

    lines.append("")
    lines.append("## Evidence Audit")
    lines.append("| Module | Automated Mapped | Auto Evidence Present | Auto Evidence Missing | Spec Mapped | Spec Evidence Present | Spec Evidence Missing |")
    lines.append("|---|---:|---:|---:|---:|---:|---:|")
    for r in rows:
        lines.append(
            f"| {r['module_key']} | {r['automated_cases']} | {r['auto_evidence_present']} | {r['auto_evidence_missing']} | {r['spec_flow_cases']} | {r['spec_evidence_present']} | {r['spec_evidence_missing']} |"
        )

    lines.append("")
    lines.append("## Notes")
    lines.append("- Coverage here is ID-mapped against manual Excel case IDs.")
    lines.append("- Automated coverage uses case IDs detected in tests and reviewed IDs from verified mapping.")
    lines.append("- Spec-flow coverage uses reviewed IDs from verified mapping only (no endpoint-based auto inference).")
    lines.append("- Pass/Fail counts are derived from the latest JUnit report; xfail cases are counted as failed.")
    lines.append(f"- Strict verified-spec mode: {'enabled' if strict_verified_spec else 'disabled'}.")
    lines.append("- Net coverage de-duplicates automated/spec overlap to avoid double counting.")
    lines.append("- Only go-signal modules contribute to the top summary coverage.")
    lines.append("- Update docs/coverage/stable_modules.json when devs mark a module stable.")

    (OUT_DIR / "STABLE_COVERAGE_DASHBOARD.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    html = f"""<!doctype html>
<html lang=\"en\">
<head>
  <meta charset=\"utf-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1\" />
  <title>Stable API Coverage Dashboard</title>
  <style>
    :root {{ --bg:#f4f7f5; --card:#ffffff; --ink:#132a1f; --muted:#5f7468; --ok:#1f7a4f; --warn:#b06500; --line:#d6e1db; }}
    body {{ margin:0; font-family: 'Segoe UI', Tahoma, sans-serif; background:linear-gradient(135deg,#f4f7f5,#e9f1ec); color:var(--ink); }}
    .wrap {{ max-width:1100px; margin:30px auto; padding:0 18px; }}
    .card {{ background:var(--card); border:1px solid var(--line); border-radius:14px; padding:16px 18px; margin-bottom:14px; box-shadow:0 8px 18px rgba(0,0,0,.04); }}
    h1 {{ margin:0 0 8px; font-size:28px; }}
    .meta {{ color:var(--muted); font-size:14px; }}
    .grid {{ display:grid; grid-template-columns: repeat(4,minmax(120px,1fr)); gap:10px; }}
    .kpi {{ background:#f8fbf9; border:1px solid var(--line); border-radius:10px; padding:10px; }}
    .kpi .v {{ font-size:22px; font-weight:700; }}
    .kpi .l {{ color:var(--muted); font-size:12px; }}
    table {{ width:100%; border-collapse:collapse; font-size:14px; }}
    th,td {{ border-bottom:1px solid var(--line); padding:9px 8px; text-align:left; }}
    th {{ background:#f6faf7; }}
    .yes {{ color:var(--ok); font-weight:600; }}
    .no {{ color:var(--warn); font-weight:600; }}
  </style>
</head>
<body>
<div class=\"wrap\">
  <div class=\"card\">
    <h1>Stable API Coverage Dashboard</h1>
    <div class=\"meta\">Generated (UTC): {payload['generated_at_utc']}</div>
  </div>
  <div class=\"card\">
    <div class=\"grid\">
      <div class=\"kpi\"><div class=\"v\">{totals['go_signal_modules']}/{totals['module_count']}</div><div class=\"l\">Go-signal modules</div></div>
      <div class=\"kpi\"><div class=\"v\">{totals['manual_cases']}</div><div class=\"l\">Manual cases (go-signal)</div></div>
      <div class=\"kpi\"><div class=\"v\">{totals['automated_cases']}</div><div class=\"l\">Automated cases</div></div>
      <div class=\"kpi\"><div class=\"v\">{totals['coverage_percent']}%</div><div class=\"l\">Coverage</div></div>
    </div>
        <div class=\"meta\" style=\"margin-top:10px\">JUnit source: {payload.get('junit_source') or 'not found'} | Passed (go-signal): {totals['passed_cases']} | Failed (go-signal): {totals['failed_cases']}</div>
  </div>
  <div class=\"card\">
    <table>
                        <thead><tr><th>Module</th><th>Go Signal</th><th>Manual</th><th>Automated</th><th>Spec Flow</th><th>Overlap</th><th>Net Covered</th><th>Pending</th><th>Passed</th><th>Failed</th><th>Coverage %</th><th>Test Files</th><th>Schema Files</th></tr></thead>
      <tbody>
                                {''.join([f"<tr><td>{r['module_key']}</td><td class='{'yes' if r['go_signal'] else 'no'}'>{'Yes' if r['go_signal'] else 'No'}</td><td>{r['manual_cases']}</td><td>{r['automated_cases']}</td><td>{r['spec_flow_cases']}</td><td>{r['overlap_cases']}</td><td>{r['net_covered_cases']}</td><td>{r['pending_cases']}</td><td>{r['passed_cases']}</td><td>{r['failed_cases']}</td><td>{r['coverage_percent']}</td><td>{r['test_files']}</td><td>{r['schema_files']}</td></tr>" for r in rows])}
      </tbody>
    </table>
  </div>
    <div class=\"card\">
        <table>
            <thead><tr><th>Module</th><th>Automated Mapped</th><th>Auto Evidence Present</th><th>Auto Evidence Missing</th><th>Spec Mapped</th><th>Spec Evidence Present</th><th>Spec Evidence Missing</th></tr></thead>
            <tbody>
                {''.join([f"<tr><td>{r['module_key']}</td><td>{r['automated_cases']}</td><td>{r['auto_evidence_present']}</td><td>{r['auto_evidence_missing']}</td><td>{r['spec_flow_cases']}</td><td>{r['spec_evidence_present']}</td><td>{r['spec_evidence_missing']}</td></tr>" for r in rows])}
            </tbody>
        </table>
    </div>
</div>
</body>
</html>
"""
    (OUT_DIR / "STABLE_COVERAGE_DASHBOARD.html").write_text(html, encoding="utf-8")

    print("Generated docs/coverage dashboard artifacts.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate stable module coverage dashboard artifacts.")
    parser.add_argument(
        "--strict-verified-spec",
        action="store_true",
        help="Fail if any case is marked spec-covered without explicit verified mapping.",
    )
    args = parser.parse_args()
    main(strict_verified_spec=args.strict_verified_spec)
