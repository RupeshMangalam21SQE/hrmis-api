from __future__ import annotations

import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Dict

from openpyxl import load_workbook


WORKBOOK_ENV = "API_REGRESSION_XLSX_PATH"
DEFAULT_WORKBOOK = "tests/test_data/HRMIS_API-Test cases .xlsx"
ID_RE = re.compile(r"^[A-Z]{2,5}-[A-Z]{1,3}-\d+$")
STATUS_RE = re.compile(r"\b([1-5]\d\d)\b")


def _root_dir() -> Path:
    return Path(__file__).resolve().parents[2]


def workbook_path() -> Path:
    p = os.getenv(WORKBOOK_ENV, "").strip()
    return Path(p) if p else _root_dir() / DEFAULT_WORKBOOK


def _clean(v) -> str:
    return "" if v is None else str(v).strip()


def _infer_role_hint(pre: str) -> str:
    s = pre.lower()
    if "superadmin" in s:
        return "superadmin"
    if "normal user" in s or "employee" in s:
        return "employee"
    if "hr" in s:
        return "hr"
    return ""


def _infer_auth_mode(pre: str, expected: str) -> str:
    p = pre.lower()
    e = expected.lower()
    if "without authentication" in e or "missing authentication" in e or "unauthorized" in e:
        return "unauthenticated"
    if "valid authentication token" in p or "authorization: bearer" in p:
        return "authenticated"
    return "unknown"


def _extract_status_codes(expected: str) -> list[int]:
    codes = []
    for m in STATUS_RE.findall(expected):
        code = int(m)
        if code not in codes:
            codes.append(code)
    return codes


def _split_endpoints(endpoint: str) -> list[str]:
    raw = endpoint.replace("\r", "\n")
    parts = []
    for piece in raw.split("\n"):
        p = piece.strip()
        if not p:
            continue
        if p == "&":
            continue
        if " " in p and p.startswith("/"):
            p = p.split()[0]
        parts.append(p)
    return parts


@lru_cache(maxsize=1)
def load_sheet_cases(sheet_name: str) -> Dict[str, dict]:
    wb_path = workbook_path()
    if not wb_path.exists():
        raise FileNotFoundError(f"Workbook not found: {wb_path}")

    wb = load_workbook(wb_path, data_only=True)
    ws = wb[sheet_name]

    header = [_clean(v).lower() for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    by_id: Dict[str, dict] = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        vals = [_clean(v) for v in row]
        case_id = vals[0]
        if not ID_RE.match(case_id):
            continue

        rec = {h: v for h, v in zip(header, vals) if h}
        pre_req = rec.get("pre-requisite", rec.get("pre-requisite ", rec.get("pre-requisite", "")))
        expected = rec.get("expected result", rec.get(" expected result ", ""))
        automation_status = rec.get("automation status", "")
        by_id[case_id] = {
            "case_id": case_id,
            "priority": rec.get("priority", ""),
            "description": rec.get("test description", ""),
            "method": rec.get("http method", ""),
            "endpoint": rec.get("api endpoint", ""),
            "endpoint_variants": _split_endpoints(rec.get("api endpoint", "")),
            "request_body": rec.get("request body", ""),
            "steps": rec.get("test steps", ""),
            "pre_requisite": pre_req,
            "expected": expected,
            "expected_status_codes": _extract_status_codes(expected),
            "automation_status": automation_status,
            "role_hint": _infer_role_hint(pre_req),
            "auth_mode": _infer_auth_mode(pre_req, expected),
        }

    return by_id
