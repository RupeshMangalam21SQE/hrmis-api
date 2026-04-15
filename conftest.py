# conftest.py
import os
import re
from io import BytesIO
from datetime import datetime, timedelta
from pathlib import Path
import pytest
from collections.abc import Generator
from typing import Dict, Optional, Union
from _pytest.fixtures import FixtureLookupError
from playwright.sync_api import Playwright, APIRequestContext, sync_playwright
from openpyxl import load_workbook, Workbook

# ---------- CLI and role selection ----------

def pytest_addoption(parser):
    parser.addoption(
        "--role",
        action="store",
        default=os.getenv("DEFAULT_ROLE", "superadmin"),
        help="Default role to run tests under (superadmin, hr, employee, l1, l2, l3, store)",
    )
    parser.addoption(
        "--report-dir",
        action="store",
        default=None,
        help="Directory where test reports are written (default: reports/<timestamp>)",
    )


def _resolve_report_dir(config) -> Path:
    user_dir = config.getoption("--report-dir") or os.getenv("REPORT_DIR")
    if user_dir:
        return Path(user_dir)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path("reports") / ts


def pytest_configure(config):
    report_dir = _resolve_report_dir(config)
    report_dir.mkdir(parents=True, exist_ok=True)
    setattr(config, "_report_dir", report_dir)
    setattr(config, "_xfail_records", [])

    # Auto-enable JUnit and HTML reports unless user already provided explicit paths.
    if hasattr(config.option, "xmlpath") and not getattr(config.option, "xmlpath", None):
        config.option.xmlpath = str(report_dir / "junit.xml")
    if hasattr(config.option, "junitxml") and not getattr(config.option, "junitxml", None):
        config.option.junitxml = str(report_dir / "junit.xml")

    if hasattr(config.option, "htmlpath") and not getattr(config.option, "htmlpath", None):
        config.option.htmlpath = str(report_dir / "report.html")
        if hasattr(config.option, "self_contained_html"):
            config.option.self_contained_html = True


def pytest_terminal_summary(terminalreporter, exitstatus, config):
    report_dir = getattr(config, "_report_dir", None)
    if not report_dir:
        return

    stats = terminalreporter.stats
    summary_lines = [
        f"exit_status: {exitstatus}",
        f"passed: {len(stats.get('passed', []))}",
        f"failed: {len(stats.get('failed', []))}",
        f"xfailed: {len(stats.get('xfailed', []))}",
        f"xpassed: {len(stats.get('xpassed', []))}",
        f"skipped: {len(stats.get('skipped', []))}",
        f"errors: {len(stats.get('error', []))}",
    ]
    (report_dir / "summary.txt").write_text("\n".join(summary_lines) + "\n", encoding="utf-8")
    _write_xfail_defect_excel(config, report_dir)


@pytest.hookimpl(hookwrapper=True)
def pytest_runtest_makereport(item, call):
    outcome = yield
    report = outcome.get_result()

    # XFAIL shows up as skipped with a wasxfail reason in call phase.
    if report.when != "call" or report.outcome != "skipped" or not hasattr(report, "wasxfail"):
        return

    reason = str(getattr(report, "wasxfail", ""))
    expected, actual = _extract_expected_actual(reason)

    case_match = re.search(r"\[([A-Z]{2,5}-[A-Z]{1,3}-\d+)\]", report.nodeid)
    if not case_match:
        case_match = re.search(r"\b([A-Z]{2,5}-[A-Z]{1,3}-\d+)\b", reason)
    defect_match = re.search(r"\[([A-Z]{2,6}-DEF-\d{3})\]", reason)

    records = getattr(item.config, "_xfail_records", [])
    records.append(
        {
            "nodeid": report.nodeid,
            "module": report.nodeid.split("::")[0],
            "test": report.nodeid.split("::")[-1],
            "case_id": case_match.group(1) if case_match else "",
            "defect_id": defect_match.group(1) if defect_match else "",
            "expected": expected,
            "actual": actual,
            "reason": reason,
        }
    )
    setattr(item.config, "_xfail_records", records)


def _extract_expected_actual(reason: str) -> tuple[str, str]:
    text = " ".join((reason or "").split())

    # Preferred shape from test helper:
    # [DEF-ID] expected=...; actual=...; note=...
    m = re.search(r"expected\s*=\s*(.+?)\s*;\s*actual\s*=\s*(.+?)(?:\s*;\s*note\s*=|$)", text, re.IGNORECASE)
    if m:
        return m.group(1).strip(), m.group(2).strip()

    # Generic shape:
    # expected X, got Y
    m = re.search(r"expected\s+(.+?)\s*,\s*got\s+(.+?)(?:\s*::|$)", text, re.IGNORECASE)
    if m:
        return m.group(1).strip(), m.group(2).strip()

    return "See xfail reason", text


def _write_xfail_defect_excel(config, report_dir: Path) -> None:
    records = getattr(config, "_xfail_records", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "xfail_defects"
    ws.append(
        [
            "run_timestamp",
            "module",
            "test",
            "case_id",
            "defect_id",
            "expected_outcome",
            "actual_outcome",
            "xfail_reason",
            "nodeid",
        ]
    )

    run_ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for rec in records:
        ws.append(
            [
                run_ts,
                rec.get("module", ""),
                rec.get("test", ""),
                rec.get("case_id", ""),
                rec.get("defect_id", ""),
                rec.get("expected", ""),
                rec.get("actual", ""),
                rec.get("reason", ""),
                rec.get("nodeid", ""),
            ]
        )

    report_file = report_dir / "xfail_defect_mapping.xlsx"
    wb.save(report_file)

    shared_file = Path("tests") / "test_data" / "xfail_defect_mapping.xlsx"
    shared_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(shared_file)

@pytest.fixture(scope="session")
def default_role(request) -> str:
    return request.config.getoption("--role")

@pytest.fixture
def ctx(request, default_role):
    """
    Generic request context:
    - Uses @pytest.mark.role("<name>") on the test if present.
    - Else uses --role or DEFAULT_ROLE=superadmin.
    """
    role_marker = request.node.get_closest_marker("role")
    role = role_marker.args[0] if role_marker and role_marker.args else default_role
    fixture_name = f"api_{role}"
    try:
        return request.getfixturevalue(fixture_name)
    except FixtureLookupError:
        pytest.fail(f"Unknown role '{role}': missing fixture '{fixture_name}'")

# ---------- URL helper ----------

def _p(path: str) -> str:
    """Join API_PREFIX with a relative path."""
    prefix = os.getenv("API_PREFIX", "HRMBackendTest").strip("/")
    return f"/{prefix}/{path.lstrip('/')}"

# ---------- Playwright bootstrap (no plugin required) ----------

@pytest.fixture(scope="session")
def playwright() -> Generator[Playwright, None, None]:
    with sync_playwright() as p:
        yield p

@pytest.fixture(scope="session")
def base_url() -> str:
    return os.getenv("BASE_URL", "https://topuptalent.com")

@pytest.fixture(scope="session")
def unauth_ctx(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    # Global default API timeout (ms)
    timeout_ms = int(os.getenv("API_TIMEOUT_MS", "30000"))
    ctx = playwright.request.new_context(base_url=base_url, timeout=timeout_ms, ignore_https_errors=True)
    yield ctx
    ctx.dispose()

# ---------- Auth helpers (JSON by default + configurable) ----------

def _signin(ctx: APIRequestContext, login_path: str, email: Optional[str], password: Optional[str]) -> str:
    """
    Sign in and return bearer token.

    SIGNIN_MODE:
      - 'data' (default): JSON body via data={} (Content-Type: application/json)
      - 'form'          : application/x-www-form-urlencoded via form={}
    Timeouts:
      - SIGNIN_TIMEOUT_MS (overrides API_TIMEOUT_MS for signin only)
    """
    assert email and password, "Environment variables for user and password are not set"
    mode = os.getenv("SIGNIN_MODE", "data").lower()
    timeout_ms = int(os.getenv("SIGNIN_TIMEOUT_MS", os.getenv("API_TIMEOUT_MS", "30000")))

    if mode == "form":
        resp = ctx.post(login_path, form={"email": email, "password": password}, timeout=timeout_ms)
    else:
        # Default to JSON body to match collections
        resp = ctx.post(
            login_path,
            data={"email": email, "password": password},
            timeout=timeout_ms,
        )

    assert resp.ok, f"Auth failed for {email}: {resp.status} :: {resp.text()}"
    body = resp.json()
    token = body.get("accessToken") or body.get("token") or body.get("access_token")
    assert token, f"No token in response: {body}"
    return token

def _auth_ctx(playwright: Playwright, base_url: str, email_env: str, pass_env: str) -> APIRequestContext:
    timeout_ms = int(os.getenv("API_TIMEOUT_MS", "30000"))
    bootstrap = playwright.request.new_context(base_url=base_url, timeout=timeout_ms, ignore_https_errors=True)
    login_path = _p("api/auth/signin")
    token = _signin(bootstrap, login_path, os.getenv(email_env), os.getenv(pass_env))
    bootstrap.dispose()
    return playwright.request.new_context(
        base_url=base_url,
        timeout=timeout_ms,
        ignore_https_errors=True,
        extra_http_headers={"Authorization": f"Bearer {token}"}
    )

# ---------- Role-scoped APIRequestContext fixtures ----------

@pytest.fixture(scope="session")
def api_hr(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    ctx = _auth_ctx(playwright, base_url, "HR_USER", "HR_PASS")
    yield ctx
    ctx.dispose()

@pytest.fixture(scope="session")
def api_employee(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    ctx = _auth_ctx(playwright, base_url, "EMPLOYEE_USER", "EMPLOYEE_PASS")
    yield ctx
    ctx.dispose()

@pytest.fixture(scope="session")
def api_superadmin(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    ctx = _auth_ctx(playwright, base_url, "SUPERADMIN_USER", "SUPERADMIN_PASS")
    yield ctx
    ctx.dispose()

@pytest.fixture(scope="session")
def api_l1(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    ctx = _auth_ctx(playwright, base_url, "ASSET_L1_USER", "ASSET_L1_PASS")
    yield ctx
    ctx.dispose()

@pytest.fixture(scope="session")
def api_l2(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    ctx = _auth_ctx(playwright, base_url, "ASSET_L2_USER", "ASSET_L2_PASS")
    yield ctx
    ctx.dispose()

@pytest.fixture(scope="session")
def api_l3(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    ctx = _auth_ctx(playwright, base_url, "ASSET_L3_USER", "ASSET_L3_PASS")
    yield ctx
    ctx.dispose()

@pytest.fixture(scope="session")
def api_store(playwright: Playwright, base_url: str) -> Generator[APIRequestContext, None, None]:
    ctx = _auth_ctx(playwright, base_url, "ASSET_STORE_USER", "ASSET_STORE_PASS")
    yield ctx
    ctx.dispose()

@pytest.fixture(scope="session")
def request_ctx(unauth_ctx: APIRequestContext) -> APIRequestContext:
    return unauth_ctx

# ---------- Identity resolver for EMPLOYEE (optional) ----------

def _clean_code(val: Optional[str]) -> Optional[str]:
    """Treat empty or placeholder values like '<...>' as unset."""
    if not val:
        return None
    s = val.strip()
    if not s or re.match(r"^<.*>$", s):
        return None
    return s

def _clean_id(val: Optional[str]) -> Optional[int]:
    """Return int if numeric; otherwise unset."""
    if not val:
        return None
    s = val.strip()
    return int(s) if s.isdigit() else None

@pytest.fixture(scope="session")
def identity_employee(api_employee: APIRequestContext) -> Dict[str, Union[str, int, None]]:
    """
    Resolve logged-in employee identity:
    - Prefer EMPLOYEE_CODE / EMPLOYEE_ID if set to real values (not '<...>').
    - Else infer from the employee's own asset list (first page).
    """
    code = _clean_code(os.getenv("EMPLOYEE_CODE"))
    emp_id = _clean_id(os.getenv("EMPLOYEE_ID"))

    if code or emp_id is not None:
        return {"employeeCode": code, "employeeId": emp_id, "source": "env"}

    r = api_employee.get(_p("assest/assestRequestList?pageSize=5&page=1"))
    if r.ok:
        body = r.json()
        items = body["data"] if isinstance(body, dict) and "data" in body else (body if isinstance(body, list) else [])
        if items:
            first = items[0]
            emp_id = first.get("employeeId")
            name = first.get("employeeName")
            code = first.get("employeeCode") or first.get("empCode")
            return {"employeeCode": code, "employeeId": emp_id, "employeeName": name, "source": "inferred"}

    return {"employeeCode": None, "employeeId": None, "source": "unknown"}


# ---------- Announcement test data fixtures ----------

def _to_iso_no_seconds(value: datetime) -> str:
    return value.replace(second=0, microsecond=0).isoformat()


@pytest.fixture(scope="session")
def valid_announcement_data() -> Dict[str, str]:
    """Canonical valid payload shared by announcement tests."""
    start = datetime.now() + timedelta(days=2)
    end = start + timedelta(hours=2)
    return {
        "title": "Quarterly Policy Update",
        "description": "Updated HR policy briefing for all teams.",
        "eventType": "Classroom Session",
        "presentedBy": "HR Team",
        "startDateTime": _to_iso_no_seconds(start),
        "endDateTime": _to_iso_no_seconds(end),
        "venue": "Cabin 02",
        "mode": "Offline",
    }


@pytest.fixture(scope="session")
def invalid_date_announcement_data(valid_announcement_data: Dict[str, str]) -> Dict[str, str]:
    """Invalid payload with end time earlier than start time."""
    start = datetime.now() + timedelta(days=2)
    end = start - timedelta(hours=1)
    return {
        **valid_announcement_data,
        "startDateTime": _to_iso_no_seconds(start),
        "endDateTime": _to_iso_no_seconds(end),
    }


@pytest.fixture(scope="session")
def sample_announcement_file() -> Dict[str, Union[str, bytes]]:
    """Playwright multipart file payload for a valid PDF attachment."""
    file_path = Path(__file__).parent / "tests" / "test_data" / "sample_announcement.pdf"
    return {
        "name": "sample_announcement.pdf",
        "mimeType": "application/pdf",
        "buffer": file_path.read_bytes(),
    }


@pytest.fixture(scope="session")
def invalid_announcement_file() -> Dict[str, Union[str, bytes]]:
    """Playwright multipart file payload for invalid file type checks."""
    file_path = Path(__file__).parent / "tests" / "test_data" / "invalid_file.exe"
    return {
        "name": "invalid_file.exe",
        "mimeType": "application/x-msdownload",
        "buffer": file_path.read_bytes(),
    }


# ---------- Salary Management test data fixtures ----------

@pytest.fixture(scope="session")
def valid_salary_excel_file() -> Dict[str, Union[str, bytes]]:
    file_path = Path(__file__).parent / "tests" / "test_data" / "salary_valid.xlsx"
    raw = file_path.read_bytes()

    # Backend parser expects text-like values for some columns and crashes on numeric cells.
    wb = load_workbook(BytesIO(raw))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, bool) or cell.value is None:
                    continue
                if isinstance(cell.value, (int, float)):
                    if isinstance(cell.value, float) and cell.value.is_integer():
                        cell.value = str(int(cell.value))
                    else:
                        cell.value = str(cell.value)

    buf = BytesIO()
    wb.save(buf)
    normalized = buf.getvalue()

    return {
        "name": file_path.name,
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "buffer": normalized,
    }


@pytest.fixture(scope="session")
def invalid_salary_text_file() -> Dict[str, Union[str, bytes]]:
    return {
        "name": "salary_invalid.txt",
        "mimeType": "text/plain",
        "buffer": b"not an excel file",
    }


@pytest.fixture(scope="session")
def corrupted_salary_excel_file() -> Dict[str, Union[str, bytes]]:
    return {
        "name": "salary_corrupted.xlsx",
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "buffer": b"corrupted-xlsx-content",
    }


@pytest.fixture(scope="session")
def invalid_salary_structure_excel_file() -> Dict[str, Union[str, bytes]]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Salary"
    ws.append(["wrongColumnA", "wrongColumnB", "wrongColumnC"])
    ws.append(["abc", "def", "ghi"])
    ws.append(["1", "2", "3"])

    buf = BytesIO()
    wb.save(buf)
    return {
        "name": "salary_invalid_structure.xlsx",
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "buffer": buf.getvalue(),
    }


@pytest.fixture(scope="session")
def oversized_salary_excel_file() -> Dict[str, Union[str, bytes]]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Salary"
    ws.append(["employeeCode", "grossSalary", "month"])

    # Generate a structurally valid workbook that exceeds backend size constraints.
    for i in range(1, 30001):
        ws.append([f"EMP-{i:05d}-{i*i}", f"{10000 + i}", f"may-{i:05d}"])

    buf = BytesIO()
    wb.save(buf)
    content = buf.getvalue()
    return {
        "name": "salary_large.xlsx",
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "buffer": content,
    }
